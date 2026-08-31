#!/usr/bin/env python3
"""Generate Wafir technical-report diagram, DOCX, and print-ready HTML."""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips

ROOT = Path(__file__).resolve().parents[1]
SUB = ROOT / "docs" / "submission"
LOGO = SUB / "wafir-logo.jpg"
DIAGRAM = SUB / "y3172-pipeline.png"
DOCX_PATH = ROOT / "Wafir—Technical Report (AI Readiness Hackathon – KSA).docx"
HTML_PATH = SUB / "technical-report.html"
GREEN = (15, 118, 110)
NAVY = (15, 23, 42)
MUTED = (71, 85, 105)
LIGHT = (240, 253, 250)
WHITE = (255, 255, 255)


def ar(s: str) -> str:
    """Isolate Arabic so it does not reverse inside LTR English."""
    return f"\u2067{s}\u2069"


def draw_pipeline(path: Path) -> None:
    w, h = 1600, 340
    img = Image.new("RGB", (w, h), WHITE)
    d = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 22)
        small = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 16)
        tiny = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 14)
    except OSError:
        font = ImageFont.load_default()
        small = font
        tiny = font

    nodes = ["SRC", "C", "PP", "M", "P", "D", "SINK"]
    labels = [
        "budget / kWh",
        "collector",
        "gap-fill",
        "forecast",
        "policy",
        "notify",
        "RTL UI",
    ]
    box_w, box_h = 150, 78
    gap = 38
    start_x = 40
    y = 48
    centers = []
    for i, (name, lab) in enumerate(zip(nodes, labels)):
        x = start_x + i * (box_w + gap)
        d.rounded_rectangle([x, y, x + box_w, y + box_h], radius=14, fill=LIGHT, outline=GREEN, width=3)
        d.text((x + box_w / 2, y + 22), name, fill=NAVY, font=font, anchor="mm")
        d.text((x + box_w / 2, y + 52), lab, fill=MUTED, font=small, anchor="mm")
        centers.append((x + box_w, y + box_h / 2, x + box_w + gap))
        if i < len(nodes) - 1:
            x1 = x + box_w
            x2 = x + box_w + gap
            mid = y + box_h / 2
            d.line([(x1 + 4, mid), (x2 - 10, mid)], fill=GREEN, width=3)
            d.polygon([(x2 - 12, mid - 7), (x2 - 2, mid), (x2 - 12, mid + 7)], fill=GREEN)

    # Sandbox under PP–M
    sx, sy = start_x + 2 * (box_w + gap) + 20, 175
    sw, sh = 280, 70
    d.rounded_rectangle([sx, sy, sx + sw, sy + sh], radius=14, fill=(254, 252, 232), outline=(180, 83, 9), width=3)
    d.text((sx + sw / 2, sy + 24), "Sandbox", fill=NAVY, font=font, anchor="mm")
    d.text((sx + sw / 2, sy + 50), "forecast, no distributor alerts", fill=MUTED, font=tiny, anchor="mm")
    pp_cx = start_x + 2 * (box_w + gap) + box_w / 2
    m_cx = start_x + 3 * (box_w + gap) + box_w / 2
    d.line([(pp_cx, y + box_h), (pp_cx, sy)], fill=(180, 83, 9), width=2)
    d.line([(m_cx, y + box_h), (m_cx, sy)], fill=(180, 83, 9), width=2)
    d.text((w / 2, h - 18), "ITU-T Y.3172  ·  SRC → C → PP → M → P → D → SINK  ·  Sandbox parallel to M", fill=MUTED, font=tiny, anchor="mm")
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path, "PNG")


def set_run_font(run, size=10, bold=False, color=NAVY, name="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), "IBM Plex Sans Arabic")
    rFonts.set(qn("w:eastAsia"), name)


def add_hyperlink(paragraph, text, url, size=8):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0F766E")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), str(int(size * 2)))
    rPr.append(szCs)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def shade_header(cell, fill="0F766E"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)
    for p in cell.paragraphs:
        for run in p.runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "D1D5DB")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def compact_paragraph(p, after=40, before=0, line=220):
    pf = p.paragraph_format
    pf.space_after = Twips(after)
    pf.space_before = Twips(before)
    pf.line_spacing = line / 240
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def add_heading(doc, text, size=12):
    p = doc.add_paragraph()
    compact_paragraph(p, after=60, before=120)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=True, color=GREEN)
    return p


def add_body(doc, text, size=10, bold=False):
    p = doc.add_paragraph()
    compact_paragraph(p, after=60)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_table(doc, headers, rows, col_widths, header_size=8, cell_size=7.5):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.autofit = False
    table.allow_autofit = False
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        compact_paragraph(p, after=20, before=20)
        run = p.add_run(h)
        set_run_font(run, size=header_size, bold=True, color=WHITE)
        shade_header(cell)
        set_cell_border(cell)
        cell.width = Inches(col_widths[i])
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            p = cell.paragraphs[0]
            compact_paragraph(p, after=20, before=20)
            if isinstance(val, tuple):
                label, url = val
                if url:
                    add_hyperlink(p, label, url, size=cell_size)
                else:
                    run = p.add_run(label)
                    set_run_font(run, size=cell_size)
            else:
                run = p.add_run(str(val))
                set_run_font(run, size=cell_size)
            set_cell_border(cell)
            if r_i % 2 == 1:
                tc = cell._tc
                tcPr = tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "F8FAFC")
                shd.set(qn("w:val"), "clear")
                tcPr.append(shd)
            cell.width = Inches(col_widths[c_i])
    spacer = doc.add_paragraph()
    compact_paragraph(spacer, after=40, before=0)
    return table


def set_narrow_margins(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.5)


def build_docx():
    doc = Document()
    set_narrow_margins(doc.sections[0])
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    # Header with logo
    if LOGO.exists():
        p = doc.add_paragraph()
        compact_paragraph(p, after=20)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run()
        run.add_picture(str(LOGO), width=Inches(1.35))

    title = doc.add_paragraph()
    compact_paragraph(title, after=40)
    r = title.add_run("Wafir — Technical Report (AI Readiness Hackathon – KSA)")
    set_run_font(r, size=14, bold=True)

    meta = [
        ("Team name: ", "wafir team"),
        ("Track: ", "Finance (household energy budgeting / FinTech)"),
        ("Solution name: ", "Wafir — Proactive Bill Prediction and Budget Planning"),
        ("Contact / Org: ", "King Faisal University — Budget Planning"),
        ("Designation: ", "Students"),
    ]
    for k, v in meta:
        p = doc.add_paragraph()
        compact_paragraph(p, after=20)
        rk = p.add_run(k)
        set_run_font(rk, size=9, bold=True)
        rv = p.add_run(v)
        set_run_font(rv, size=9)

    members_heading = doc.add_paragraph()
    compact_paragraph(members_heading, after=20)
    mh = members_heading.add_run("Team Members")
    set_run_font(mh, size=9, bold=True)

    members = [
        "Mohammed Nadher Aboalrejal — Mentor — aboalrejal.ai@gmail.com — Technical Lead, System Architect, Technical Report Author",
        "Fatima Alsultan (KFU, Chemical Engineering, 2nd Year) — fatima.alsultan2105@gmail.com — Document research, concept ideation, policy gap analysis, report structuring",
        "Shahad Alsultan (KFU, Civil Engineering, 2nd Year) — shahadalsultan2026@outlook.com — Video production, report drafting (Introduction and Nodes)",
        "Jorry Alfalah (KFU, Electrical Engineering, 2nd Year) — Jurryraed90@hotmail.com — App UI design, logo design",
        "Noor Alshammari (KFU, Chemical Engineering, 2nd Year) — noornaser.sh1@gmail.com — App UI templates, logo design",
    ]
    for member in members:
        p = doc.add_paragraph()
        compact_paragraph(p, after=12)
        run = p.add_run("• " + member)
        set_run_font(run, size=8.5)

    cap = doc.add_paragraph()
    compact_paragraph(cap, after=40)
    rc = cap.add_run(
        "Cap: ≤5 pages. Demo telemetry is synthetic (hackathon use only); Knowledge Base documents are authentic public sources. This report describes only capabilities present in the repository (demo path + optional Supabase). Limitations are stated explicitly."
    )
    set_run_font(rc, size=8.5, color=MUTED)

    res_h = doc.add_paragraph()
    compact_paragraph(res_h, after=20)
    rh = res_h.add_run("Project resources")
    set_run_font(rh, size=9, bold=True)

    for label, url in [
        ("Live Demo: ", "https://wafier.aboalrejal.com/"),
        ("GitHub: ", "https://github.com/aboalrejal-ai/wafier"),
        ("Demo video: ", "https://www.youtube.com/watch?v=bRYwjbxs9t4"),
    ]:
        p = doc.add_paragraph()
        compact_paragraph(p, after=12)
        run = p.add_run("• " + label)
        set_run_font(run, size=9)
        add_hyperlink(p, url, url, size=9)

    add_heading(doc, "1. Introduction")
    add_body(
        doc,
        "Households in Saudi Arabia often discover electricity overspend only when the monthly bill arrives. Peak summer cooling makes this worse. Wafir is an Arabic RTL personal FinTech app that forecasts end-of-month SAR spend from simulated meter-like kWh plus weather context, compares the forecast to a user budget, and raises graduated alerts.",
    )
    add_body(
        doc,
        "The pipeline follows ITU-T Y.3172 (SRC → C → PP → M → P → D → SINK, plus Sandbox). Policy grounding uses a verified Knowledge Base of Saudi and ITU public sources (knowledge-base.json, 42 VERIFIED records + 1 ISO benchmark), compiled from three independent deep-research passes in docs/research/. Framework names follow ITU AI Readiness Report 2.0 (kb/framework/dimensions.json).",
    )
    add_body(
        doc,
        "Honest limit: meter ingest is simulated; the forecast is rule-based seasonal projection, not a trained neural model.",
    )

    add_heading(doc, "2. Use case and gaps in existing solutions")
    add_body(
        doc,
        "Problem. Families lack proactive, explainable energy-budget control tied to local tariffs and heatwaves. Beneficiaries: Saudi households managing a monthly electricity budget; secondary value for energy/privacy policy readers via the gap matrix.",
    )
    add_table(
        doc,
        ["Gap in existing tools", "How Wafir responds (implemented)"],
        [
            ["Late bill shock", "Seasonal forecast vs budget; Level-2 warning if predicted SAR > budget or temperature ≥40°C"],
            ["Opaque advice", "RAG over verified chunks with citation URLs; if nothing matches: Arabic INSUFFICIENT_EVIDENCE, no invented law"],
            ["No privacy gate", "PDPL consent screen before processing; preprocessor pseudonymizes household id before ML export"],
            ["Ads from consumption", "Deterministic policy node blocks targeted_ads_from_consumption (PDPL-ADS-001)"],
            ["No orchestration evidence", "Demo path SRC→C→PP→M→P→D with audit log; reproducible via pnpm demo"],
        ],
        [2.2, 4.9],
    )

    add_heading(doc, "3. Y.3172 architecture (mapped documents)")
    add_body(
        doc,
        "Sandbox is a Y.3172 supporting component (parallel validation). MLFO is seasonal-profile switching inside M, not an eighth official node. Y.3172 ↔ Readiness (ITU summary): D5 workflows, D10 policies, D13 infrastructure.",
        size=9,
    )
    if DIAGRAM.exists():
        p = doc.add_paragraph()
        compact_paragraph(p, after=40)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(DIAGRAM), width=Inches(7.0))

    add_table(
        doc,
        ["Node", "Role in Wafir", "Supporting document"],
        [
            ["SRC", "Budget 500 SAR, simulated kWh, weather (OpenWeather if key set, else simulated)", ("SERA residential tariff", "https://www.sera.gov.sa/en/consumer/electric-tariff/electric-tariff-categories/consumption-tariff")],
            ["C", "Demo collector / heatwave aggregation (simulateMeterReading). No consumer AMI API — GAP-01", ("SEC data sharing", "https://www.se.com.sa/en/Open-Data/Data-Sharing/")],
            ["PP", "Gap-fill missing daily kWh; hash household id (pseudonymization, not full anonymization)", ("PDPL destruction / pseudonymisation", "https://dgp.sdaia.gov.sa/wps/portal/pdp/knowledgecenter/details/PersonalDataDestruction/")],
            ["M + MLFO", "Rule-based SAR forecast; summer profile at ≥38°C (method reference, not trained weights)", ("MDPI KSA energy forecasting", "https://www.mdpi.com/1996-1073/16/4/2035")],
            ["Sandbox", "isSandbox: forecast without distributor alerts", ("ITU-T Y.3172", "https://www.itu.int/rec/T-REC-Y.3172-201906-I")],
            ["P", "Budget L1/L1b/L2 + ads guard (no LLM in the decision)", ("PDPL official text + PDPL-ADS-001", "https://sdaia.gov.sa/en/SDAIA/about/Documents/Personal%20Data%20English%20V2-23April2023-%20Reviewed-.pdf")],
            ["D", "In-app + Web/Local notifications (FCM needs keys)", ("SDAIA AI Ethics (human oversight)", "https://dgp.sdaia.gov.sa/wps/wcm/connect/4c56ed1c-1b82-447d-ac29-638f5f99c12e/ai-principles-EN.pdf")],
            ["SINK", "Arabic RTL React UI. SAMA consumer-protection analog; not Open Banking integration", ("SAMA consumer protection", "https://rulebook.sama.gov.sa/en/financial-consumer-protection-principles-and-rules")],
        ],
        [1.1, 3.4, 2.6],
    )

    add_heading(doc, "4. ITU AI Readiness 2.0 — six factors and selected dimensions")
    add_body(
        doc,
        "Official names: kb/framework/dimensions.json (ITU AI Readiness Report 2.0). PDPL / purpose limitation maps to D10, not D8.",
        size=9,
    )
    add_table(
        doc,
        ["Factor", "Wafir evidence"],
        [
            ["Data", "Public Saudi/ITU documents in knowledge-base.json (not private household dumps)"],
            ["Research", "Three deep-research passes in docs/research/ (ChatGPT, Gemini, Perplexity), then human merge"],
            ["Deployment Support", "Vite SPA on Hostinger; optional Supabase; Capacitor wrapper"],
            ["Standards", "Y.3172 node mapping + ITU-T Y.3172 citation"],
            ["Open Source and Code", "Public GitHub https://github.com/aboalrejal-ai/wafier"],
            ["Sandbox Environments", "isSandbox heatwave path — forecast without pushing alerts"],
        ],
        [2.0, 5.1],
    )
    add_table(
        doc,
        ["ID", "Official ITU dimension", "Wafir evidence"],
        [
            ["D5", "Level of Integration of AI in Workflows", "Pipeline audit, graduated alerts, pnpm demo"],
            ["D6", "Human Interface", "Arabic RTL UI + AI assistant with citations"],
            ["D7", "Strategy Alignment", "KB NSDAI-001 (Vision 2030 / SDAIA NSDAI PDF). Use-case aligned to national strategy; we do not claim to implement NSDAI"],
            ["D8", "Collaboration with AI", "HITL: 2-hour alert snooze on Profile; user question shapes RAG"],
            ["D10", "AI & Policies", "Six verdicts; SC-03 policy sandbox; GAP-01..06; PDPL consent + purpose limitation"],
            ["D11", "AI for Inclusion", "Arabic-first interface; local SERA/KAPSARC context"],
            ["D13", "Digital Infrastructure", "Hosted SPA; optional Supabase; simulated meter as Y.3172 SRC stand-in"],
        ],
        [0.6, 2.6, 3.9],
    )

    add_heading(doc, "5. Three evaluation scenarios + Policy Gap Matrix")
    add_body(
        doc,
        "Reproducible proof: pnpm demo (or node scripts/run-demo.mjs). Demo data are synthetic.",
        size=9,
    )

    add_body(doc, "SC-01 — Compliant RAG (scenarios/sc-01-compliant-rag.json)", size=10, bold=True)
    add_body(
        doc,
        "Scenario. User asks in the AI Assistant: "
        + ar("كيف أوفر في فاتورة الكهرباء؟")
        + "  Step 1: The household submits a normal saving question. Context is synthetic.  Step 2: Wafir retrieves verified KB chunks by keyword matching over curated records (not a live web search). If no chunk matches, the assistant returns "
        + ar("لا توجد أدلة كافية…")
        + " (INSUFFICIENT_EVIDENCE) and invents no URLs.  Step 3: For this query the grounded reply gives qualitative saving tips (runtime, efficient AC / SASO 2663, insulation, LED, night-time loads) with citation URLs. Expected verdict: COMPLIANT. Tips are not a guaranteed SAR saving.",
        size=9,
    )

    add_body(doc, "SC-02 — Operational failure (scenarios/sc-02-pp-gap-fill.json)", size=10, bold=True)
    add_body(
        doc,
        "Scenario. In a 7-day kWh series, day 4 is missing (kwh: null); other days are 29.67.  Step 1: Collector/preprocessor sees an incomplete series (simulated meter gap, not a live AMI outage).  Step 2: The missing day is filled from neighbouring values (demo expected fill = 29.67). Household id is pseudonymized for export.  Step 3: Forecast and dashboard continue; the audit log records the gap-fill; the app does not crash. Expected: graceful degradation. The filled value is not claimed as a real meter reading.",
        size=9,
    )

    add_body(doc, "SC-03 — Controversy / targeted ads (scenarios/sc-03-ads-controversy.json)", size=10, bold=True)
    add_body(
        doc,
        "Scenario. A dummy provider requests targeted_ads_from_consumption.  Step 1: The request arrives as a policy check (About → controversy scenario), not as a live ad SDK.  Step 2: Policy node evaluateKbGuardPolicy runs deterministically (no LLM).  Step 3: Result VIOLATION / BLOCK_DATA_USE / PDPL-ADS-001. In-app notice: "
        + ar("حارس سياسة KB — منع الإعلانات")
        + ". Audit stores the verdict and the PDPL knowledge-center URL.  Step 4 (corrective): Consumption data is not released for ads. This is a Wafir purpose-limitation control. PDPL Art.26 allows consented marketing of non-sensitive data; we do not claim a statutory total advertising ban. HITL snooze remains available on Profile for budget alerts.",
        size=9,
    )

    add_body(
        doc,
        "Policy Gap Matrix. Classifications are from sources we reviewed; they are not a claim that Saudi law is silent in every forum.",
        size=9,
    )
    add_table(
        doc,
        ["ID", "Type", "Gap", "Recommendation"],
        [
            ["GAP-01", "potential_gap", "No consumer-authorized AMI API", "Open-Utility style authorization (like Open Banking)"],
            ["GAP-02", "potential_gap", "AI forecast liability allocation", "Sector guidance; general civil law may apply (CIVIL-LAW-120)"],
            ["GAP-03", "potential_gap", "Energy-sector AI data rules", "PDPL + SERA PDP under explicit sector policy"],
            ["GAP-04", "potential_gap", "No mandatory energy algorithm-audit standard", "Y.3172 sandbox + ISO 42001 as a benchmark (not a certification)"],
            ["GAP-05", "potential_gap", "No machine-readable law repository", "Government structured-regulation API"],
            ["GAP-06", "ambiguity", "Consumption → ads repurposing", "Explicit PDPL/sector guidance; Wafir blocks silent reuse"],
        ],
        [0.85, 1.25, 2.4, 2.6],
    )
    add_body(doc, "GAP-01. In the sources we reviewed, we did not find a consumer-authorized AMI API comparable to Open Banking that would let Wafir pull live meter data with user consent. Wafir therefore uses simulated kWh. Recommendation: an Open-Utility style authorization framework.", size=8.5)
    add_body(doc, "GAP-02. We did not find sector-specific rules allocating liability when an energy forecast is wrong. General civil-law harmful-act principles may still apply (CIVIL-LAW-120). Recommendation: sector guidance; until then Wafir labels forecasts as estimates.", size=8.5)
    add_body(doc, "GAP-03. PDPL and SERA data-protection pages exist, but energy-AI processing (forecasting, sharing, secondary use) is not spelled out as a sector playbook in the sources we checked. Recommendation: PDPL + SERA PDP under explicit sector policy.", size=8.5)
    add_body(doc, "GAP-04. We did not find a mandatory energy-sector algorithm-audit standard for household bill models. Recommendation: use the Y.3172 sandbox pattern and ISO/IEC 42001 as a benchmark, not a claimed certification.", size=8.5)
    add_body(doc, "GAP-05. Official texts are published as HTML/PDF, not as a queryable regulation API. Wafir therefore curates verified chunks rather than scraping live law. Recommendation: a government structured-regulation API.", size=8.5)
    add_body(doc, "GAP-06. PDPL purpose limitation and Art.25–27 constrain marketing, but Art.26 still allows consented marketing of non-sensitive data. Whether electricity-consumption profiles may be reused for ads is not spelled out for this use case. Wafir blocks silent repurposing by policy (SC-03). Recommendation: explicit PDPL/sector guidance on energy-data repurposing.", size=8.5)

    add_heading(doc, "6. Knowledge base (authentic sources) and repository links")
    add_body(doc, "Full corpus: knowledge-base.json (42 VERIFIED + 1 ISO benchmark). Ten primary citations for judges:", size=9)
    add_table(
        doc,
        ["#", "Document", "Authority", "URL"],
        [
            ["1", "PDPL official EN (Art.4–6, 10–15, 25–27)", "SDAIA", ("sdaia.gov.sa PDF", "https://sdaia.gov.sa/en/SDAIA/about/Documents/Personal%20Data%20English%20V2-23April2023-%20Reviewed-.pdf")],
            ["2", "PDPL Art.26 direct marketing", "SDAIA", ("dgp.sdaia.gov.sa PDPL2", "https://dgp.sdaia.gov.sa/wps/portal/pdp/knowledgecenter/details/PDPL2/")],
            ["3", "AI Ethics Principles v1.0", "SDAIA", ("ai-principles-EN.pdf", "https://dgp.sdaia.gov.sa/wps/wcm/connect/4c56ed1c-1b82-447d-ac29-638f5f99c12e/ai-principles-EN.pdf")],
            ["4", "NSDAI (Vision 2030)", "SDAIA", ("NSDAI.pdf", "https://sdaia.gov.sa/en/SDAIA/SdaiaStrategies/Documents/NSDAI.pdf")],
            ["5", "SERA residential consumption tariff", "SERA", ("sera.gov.sa tariff", "https://www.sera.gov.sa/en/consumer/electric-tariff/electric-tariff-categories/consumption-tariff")],
            ["6", "SEC data sharing", "SEC", ("se.com.sa Open Data", "https://www.se.com.sa/en/Open-Data/Data-Sharing/")],
            ["7", "KBEAT — Building Energy Assessment Tool", "KAPSARC", ("apps.kapsarc.org/kbeat", "https://apps.kapsarc.org/appboard/kbeat/en")],
            ["8", "Financial Consumer Protection Principles (analog)", "SAMA", ("rulebook.sama.gov.sa", "https://rulebook.sama.gov.sa/en/financial-consumer-protection-principles-and-rules")],
            ["9", "ITU-T Y.3172 ML pipeline architecture", "ITU", ("itu.int Y.3172", "https://www.itu.int/rec/T-REC-Y.3172-201906-I")],
            ["10", "ISO/IEC 42001:2023 (benchmark only)", "ISO", ("iso.org 81230", "https://www.iso.org/standard/81230.html")],
        ],
        [0.4, 2.8, 1.1, 2.8],
    )
    add_table(
        doc,
        ["Deliverable", "Location"],
        [
            ["GitHub", ("https://github.com/aboalrejal-ai/wafier", "https://github.com/aboalrejal-ai/wafier")],
            ["Live demo", ("https://wafier.aboalrejal.com/", "https://wafier.aboalrejal.com/")],
            ["Demo video", ("https://www.youtube.com/watch?v=bRYwjbxs9t4", "https://www.youtube.com/watch?v=bRYwjbxs9t4")],
            ["Knowledge base", "knowledge-base.json"],
            ["Scenarios", "scenarios/"],
            ["Reproducible demo", "pnpm demo"],
        ],
        [2.0, 5.1],
    )
    add_body(doc, "Distributor: in-app + Web Notification API + Capacitor Local Notifications.", size=9)

    doc.save(DOCX_PATH)
    print("wrote", DOCX_PATH)


def build_html():
    logo_rel = "wafir-logo.jpg"
    diag_rel = "y3172-pipeline.png"
    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8"/>
<title>Wafir — Technical Report (AI Readiness Hackathon – KSA)</title>
<style>
  @page {{ size: letter; margin: 0.55in 0.65in 0.5in 0.65in; }}
  html, body {{ margin: 0; padding: 0; }}
  body {{
    font-family: Calibri, "IBM Plex Sans Arabic", "Segoe UI", sans-serif;
    font-size: 10pt;
    line-height: 1.28;
    color: #0f172a;
  }}
  h1 {{ font-size: 15pt; margin: 4px 0 8px; color: #0f172a; }}
  h2 {{ font-size: 11.5pt; color: #0f766e; margin: 11px 0 5px; border-bottom: 1px solid #ccfbf1; padding-bottom: 2px; }}
  h3 {{ font-size: 10.5pt; margin: 8px 0 3px; color: #134e4a; }}
  p {{ margin: 0 0 6px; }}
  .meta {{ font-size: 9.5pt; margin: 0 0 2px; }}
  .muted {{ color: #475569; font-size: 8.5pt; }}
  .members {{ font-size: 8.5pt; line-height: 1.3; margin: 0 0 2px; }}
  img.logo {{ height: 28px; }}
  img.diagram {{ width: 100%; height: auto; margin: 4px 0 8px; }}
  a {{ color: #0f766e; }}
  [lang="ar"], bdi[dir="rtl"] {{ unicode-bidi: isolate; direction: rtl; font-family: "IBM Plex Sans Arabic", "Geeza Pro", "Arial Unicode MS", sans-serif; }}
  table {{ width: 100%; border-collapse: collapse; margin: 4px 0 8px; font-size: 8pt; }}
  th {{ background: #0f766e; color: #fff; text-align: left; padding: 3px 5px; font-weight: 700; }}
  td {{ border: 1px solid #e5e7eb; padding: 3px 5px; vertical-align: top; }}
  tr:nth-child(even) td {{ background: #f8fafc; }}
  .gap {{ font-size: 8.5pt; margin-bottom: 5px; }}
  .step {{ margin-bottom: 7px; }}
</style>
</head>
<body>
<img class="logo" src="{logo_rel}" alt="Wafir"/>
<h1>Wafir — Technical Report (AI Readiness Hackathon – KSA)</h1>
<p class="meta"><b>Team name:</b> wafir team &nbsp;·&nbsp; <b>Track:</b> Finance (household energy budgeting / FinTech)</p>
<p class="meta"><b>Solution name:</b> Wafir — Proactive Bill Prediction and Budget Planning</p>
<p class="meta"><b>Contact / Org:</b> King Faisal University — Budget Planning &nbsp;·&nbsp; <b>Designation:</b> Students</p>
<p class="members"><b>Team Members</b></p>
<p class="members">• Mohammed Nadher Aboalrejal — Mentor — aboalrejal.ai@gmail.com — Technical Lead, System Architect, Technical Report Author</p>
<p class="members">• Fatima Alsultan (KFU, Chemical Engineering, 2nd Year) — fatima.alsultan2105@gmail.com — Document research, concept ideation, policy gap analysis, report structuring</p>
<p class="members">• Shahad Alsultan (KFU, Civil Engineering, 2nd Year) — shahadalsultan2026@outlook.com — Video production, report drafting (Introduction and Nodes)</p>
<p class="members">• Jorry Alfalah (KFU, Electrical Engineering, 2nd Year) — Jurryraed90@hotmail.com — App UI design, logo design</p>
<p class="members">• Noor Alshammari (KFU, Chemical Engineering, 2nd Year) — noornaser.sh1@gmail.com — App UI templates, logo design</p>
<p class="muted">Cap: ≤5 pages. Demo telemetry is <b>synthetic</b> (hackathon use only); Knowledge Base documents are authentic public sources. This report describes only capabilities present in the repository (demo path + optional Supabase). Limitations are stated explicitly.</p>
<p><b>Project resources</b></p>
<p>• Live Demo: <a href="https://wafier.aboalrejal.com/">https://wafier.aboalrejal.com/</a></p>
<p>• GitHub: <a href="https://github.com/aboalrejal-ai/wafier">https://github.com/aboalrejal-ai/wafier</a></p>
<p>• Demo video: <a href="https://www.youtube.com/watch?v=bRYwjbxs9t4">https://www.youtube.com/watch?v=bRYwjbxs9t4</a></p>

<h2>1. Introduction</h2>
<p>Households in Saudi Arabia often discover electricity overspend only when the monthly bill arrives. Peak summer cooling makes this worse. <b>Wafir</b> is an Arabic RTL personal FinTech app that forecasts end-of-month SAR spend from <b>simulated</b> meter-like kWh plus weather context, compares the forecast to a user budget, and raises graduated alerts.</p>
<p>The pipeline follows <b>ITU-T Y.3172</b> (SRC → C → PP → M → P → D → SINK, plus Sandbox). Policy grounding uses a verified Knowledge Base of Saudi and ITU public sources (<code>knowledge-base.json</code>, 42 VERIFIED records + 1 ISO benchmark), compiled from three independent deep-research passes in <code>docs/research/</code>. Framework names follow ITU AI Readiness Report 2.0 (<code>kb/framework/dimensions.json</code>).</p>
<p>Honest limit: meter ingest is simulated; the forecast is rule-based seasonal projection, not a trained neural model.</p>

<h2>2. Use case and gaps in existing solutions</h2>
<p><b>Problem.</b> Families lack proactive, explainable energy-budget control tied to local tariffs and heatwaves. <b>Beneficiaries:</b> Saudi households managing a monthly electricity budget; secondary value for energy/privacy policy readers via the gap matrix.</p>
<table>
<thead><tr><th>Gap in existing tools</th><th>How Wafir responds (implemented)</th></tr></thead>
<tbody>
<tr><td>Late bill shock</td><td>Seasonal forecast vs budget; Level-2 warning if predicted SAR &gt; budget or temperature ≥40°C</td></tr>
<tr><td>Opaque advice</td><td>RAG over verified chunks with citation URLs; if nothing matches: Arabic INSUFFICIENT_EVIDENCE, no invented law</td></tr>
<tr><td>No privacy gate</td><td>PDPL consent screen before processing; preprocessor pseudonymizes household id before ML export</td></tr>
<tr><td>Ads from consumption</td><td>Deterministic policy node blocks targeted_ads_from_consumption (PDPL-ADS-001)</td></tr>
<tr><td>No orchestration evidence</td><td>Demo path SRC→C→PP→M→P→D with audit log; reproducible via <code>pnpm demo</code></td></tr>
</tbody>
</table>

<h2>3. Y.3172 architecture (mapped documents)</h2>
<p>Sandbox is a Y.3172 supporting component (parallel validation). MLFO is seasonal-profile switching <b>inside M</b>, not an eighth official node. Y.3172 ↔ Readiness (ITU summary): D5 workflows, D10 policies, D13 infrastructure.</p>
<img class="diagram" src="{diag_rel}" alt="ITU-T Y.3172 pipeline"/>
<table>
<thead><tr><th>Node</th><th>Role in Wafir</th><th>Supporting document</th></tr></thead>
<tbody>
<tr><td>SRC</td><td>Budget 500 SAR, simulated kWh, weather (OpenWeather if key set, else simulated)</td><td><a href="https://www.sera.gov.sa/en/consumer/electric-tariff/electric-tariff-categories/consumption-tariff">SERA residential tariff</a></td></tr>
<tr><td>C</td><td>Demo collector / heatwave aggregation (simulateMeterReading). No consumer AMI API — GAP-01</td><td><a href="https://www.se.com.sa/en/Open-Data/Data-Sharing/">SEC data sharing</a></td></tr>
<tr><td>PP</td><td>Gap-fill missing daily kWh; hash household id (pseudonymization, not full anonymization)</td><td><a href="https://dgp.sdaia.gov.sa/wps/portal/pdp/knowledgecenter/details/PersonalDataDestruction/">PDPL destruction / pseudonymisation</a></td></tr>
<tr><td>M + MLFO</td><td>Rule-based SAR forecast; summer profile at ≥38°C (method reference, not trained weights)</td><td><a href="https://www.mdpi.com/1996-1073/16/4/2035">MDPI KSA energy forecasting</a></td></tr>
<tr><td>Sandbox</td><td>isSandbox: forecast without distributor alerts</td><td><a href="https://www.itu.int/rec/T-REC-Y.3172-201906-I">ITU-T Y.3172</a></td></tr>
<tr><td>P</td><td>Budget L1/L1b/L2 + ads guard (no LLM in the decision)</td><td><a href="https://sdaia.gov.sa/en/SDAIA/about/Documents/Personal%20Data%20English%20V2-23April2023-%20Reviewed-.pdf">PDPL official text + PDPL-ADS-001</a></td></tr>
<tr><td>D</td><td>In-app + Web/Local notifications (FCM needs keys)</td><td><a href="https://dgp.sdaia.gov.sa/wps/wcm/connect/4c56ed1c-1b82-447d-ac29-638f5f99c12e/ai-principles-EN.pdf">SDAIA AI Ethics (human oversight)</a></td></tr>
<tr><td>SINK</td><td>Arabic RTL React UI. SAMA analog; not Open Banking integration</td><td><a href="https://rulebook.sama.gov.sa/en/financial-consumer-protection-principles-and-rules">SAMA consumer protection</a></td></tr>
</tbody>
</table>

<h2>4. ITU AI Readiness 2.0 — six factors and selected dimensions</h2>
<p>Official names: <code>kb/framework/dimensions.json</code> (ITU AI Readiness Report 2.0). PDPL / purpose limitation maps to <b>D10</b>, not D8.</p>
<table>
<thead><tr><th>Factor</th><th>Wafir evidence</th></tr></thead>
<tbody>
<tr><td>Data</td><td>Public Saudi/ITU documents in knowledge-base.json (not private household dumps)</td></tr>
<tr><td>Research</td><td>Three deep-research passes in docs/research/ (ChatGPT, Gemini, Perplexity), then human merge</td></tr>
<tr><td>Deployment Support</td><td>Vite SPA on Hostinger; optional Supabase; Capacitor wrapper</td></tr>
<tr><td>Standards</td><td>Y.3172 node mapping + ITU-T Y.3172 citation</td></tr>
<tr><td>Open Source and Code</td><td>Public GitHub https://github.com/aboalrejal-ai/wafier</td></tr>
<tr><td>Sandbox Environments</td><td>isSandbox heatwave path — forecast without pushing alerts</td></tr>
</tbody>
</table>
<table>
<thead><tr><th>ID</th><th>Official ITU dimension</th><th>Wafir evidence</th></tr></thead>
<tbody>
<tr><td>D5</td><td>Level of Integration of AI in Workflows</td><td>Pipeline audit, graduated alerts, pnpm demo</td></tr>
<tr><td>D6</td><td>Human Interface</td><td>Arabic RTL UI + AI assistant with citations</td></tr>
<tr><td>D7</td><td>Strategy Alignment</td><td>KB NSDAI-001 (Vision 2030 / SDAIA NSDAI PDF). Use-case aligned to national strategy; we do <b>not</b> claim to implement NSDAI</td></tr>
<tr><td>D8</td><td>Collaboration with AI</td><td>HITL: 2-hour alert snooze on Profile; user question shapes RAG</td></tr>
<tr><td>D10</td><td>AI &amp; Policies</td><td>Six verdicts; SC-03 policy sandbox; GAP-01..06; PDPL consent + purpose limitation</td></tr>
<tr><td>D11</td><td>AI for Inclusion</td><td>Arabic-first interface; local SERA/KAPSARC context</td></tr>
<tr><td>D13</td><td>Digital Infrastructure</td><td>Hosted SPA; optional Supabase; simulated meter as Y.3172 SRC stand-in</td></tr>
</tbody>
</table>

<h2>5. Three evaluation scenarios + Policy Gap Matrix</h2>
<p>Reproducible proof: <code>pnpm demo</code> (or <code>node scripts/run-demo.mjs</code>). Demo data are synthetic.</p>

<div class="step"><h3>SC-01 — Compliant RAG (<code>scenarios/sc-01-compliant-rag.json</code>)</h3>
<p><b>Scenario.</b> User asks in the AI Assistant: <b><bdi dir="rtl" lang="ar">كيف أوفر في فاتورة الكهرباء؟</bdi></b> &nbsp; <b>Step 1.</b> The household submits a normal saving question. Context is synthetic. &nbsp; <b>Step 2.</b> Wafir retrieves verified KB chunks by keyword matching over curated records (not a live web search). If no chunk matches, the assistant returns <bdi dir="rtl" lang="ar">لا توجد أدلة كافية…</bdi> (INSUFFICIENT_EVIDENCE) and invents no URLs. &nbsp; <b>Step 3.</b> For this query the grounded reply gives qualitative saving tips (runtime, efficient AC / SASO 2663, insulation, LED, night-time loads) with citation URLs. Expected verdict: <b>COMPLIANT</b>. Tips are not a guaranteed SAR saving.</p></div>

<div class="step"><h3>SC-02 — Operational failure (<code>scenarios/sc-02-pp-gap-fill.json</code>)</h3>
<p><b>Scenario.</b> In a 7-day kWh series, <b>day 4 is missing</b> (kwh: null); other days are 29.67. &nbsp; <b>Step 1.</b> Collector/preprocessor sees an incomplete series (simulated meter gap, not a live AMI outage). &nbsp; <b>Step 2.</b> The missing day is filled from neighbouring values (demo expected fill = 29.67). Household id is pseudonymized for export. &nbsp; <b>Step 3.</b> Forecast and dashboard continue; the audit log records the gap-fill; the app does not crash. Expected: graceful degradation. The filled value is not claimed as a real meter reading.</p></div>

<div class="step"><h3>SC-03 — Controversy / targeted ads (<code>scenarios/sc-03-ads-controversy.json</code>)</h3>
<p><b>Scenario.</b> A dummy provider requests <code>targeted_ads_from_consumption</code>. &nbsp; <b>Step 1.</b> The request arrives as a policy check (About → controversy scenario), not as a live ad SDK. &nbsp; <b>Step 2.</b> Policy node evaluateKbGuardPolicy runs <b>deterministically</b> (no LLM). &nbsp; <b>Step 3.</b> Result <b>VIOLATION / BLOCK_DATA_USE / PDPL-ADS-001</b>. In-app notice: <bdi dir="rtl" lang="ar">حارس سياسة KB — منع الإعلانات</bdi>. Audit stores the verdict and the PDPL knowledge-center URL. &nbsp; <b>Step 4 (corrective).</b> Consumption data is not released for ads. This is a <b>Wafir purpose-limitation control</b>. PDPL Art.26 allows consented marketing of non-sensitive data; we do not claim a statutory total advertising ban. HITL snooze remains available on Profile for budget alerts.</p></div>

<p>Classifications below are from sources <b>we reviewed</b>; they are not a claim that Saudi law is silent in every forum.</p>
<table>
<thead><tr><th>ID</th><th>Type</th><th>Gap</th><th>Recommendation</th></tr></thead>
<tbody>
<tr><td>GAP-01</td><td>potential_gap</td><td>No consumer-authorized AMI API</td><td>Open-Utility style authorization (like Open Banking)</td></tr>
<tr><td>GAP-02</td><td>potential_gap</td><td>AI forecast liability allocation</td><td>Sector guidance; general civil law may apply (CIVIL-LAW-120)</td></tr>
<tr><td>GAP-03</td><td>potential_gap</td><td>Energy-sector AI data rules</td><td>PDPL + SERA PDP under explicit sector policy</td></tr>
<tr><td>GAP-04</td><td>potential_gap</td><td>No mandatory energy algorithm-audit standard</td><td>Y.3172 sandbox + ISO 42001 as a benchmark (not a certification)</td></tr>
<tr><td>GAP-05</td><td>potential_gap</td><td>No machine-readable law repository</td><td>Government structured-regulation API</td></tr>
<tr><td>GAP-06</td><td>ambiguity</td><td>Consumption → ads repurposing</td><td>Explicit PDPL/sector guidance; Wafir blocks silent reuse</td></tr>
</tbody>
</table>
<p class="gap"><b>GAP-01.</b> In the sources we reviewed, we did not find a consumer-authorized AMI API comparable to Open Banking that would let Wafir pull live meter data with user consent. Wafir therefore uses simulated kWh. Recommendation: an Open-Utility style authorization framework.</p>
<p class="gap"><b>GAP-02.</b> We did not find sector-specific rules allocating liability when an energy forecast is wrong. General civil-law harmful-act principles may still apply (CIVIL-LAW-120). Recommendation: sector guidance; until then Wafir labels forecasts as estimates.</p>
<p class="gap"><b>GAP-03.</b> PDPL and SERA data-protection pages exist, but energy-AI processing (forecasting, sharing, secondary use) is not spelled out as a sector playbook in the sources we checked. Recommendation: PDPL + SERA PDP under explicit sector policy.</p>
<p class="gap"><b>GAP-04.</b> We did not find a mandatory energy-sector algorithm-audit standard for household bill models. Recommendation: use the Y.3172 sandbox pattern and ISO/IEC 42001 as a benchmark, not a claimed certification.</p>
<p class="gap"><b>GAP-05.</b> Official texts are published as HTML/PDF, not as a queryable regulation API. Wafir therefore curates verified chunks rather than scraping live law. Recommendation: a government structured-regulation API.</p>
<p class="gap"><b>GAP-06.</b> PDPL purpose limitation and Art.25–27 constrain marketing, but Art.26 still allows consented marketing of non-sensitive data. Whether electricity-consumption profiles may be reused for ads is not spelled out for this use case. Wafir blocks silent repurposing by policy (SC-03). Recommendation: explicit PDPL/sector guidance on energy-data repurposing.</p>

<h2>6. Knowledge base (authentic sources) and repository links</h2>
<p>Full corpus: <code>knowledge-base.json</code> (42 VERIFIED + 1 ISO benchmark). Ten primary citations for judges:</p>
<table>
<thead><tr><th>#</th><th>Document</th><th>Authority</th><th>URL</th></tr></thead>
<tbody>
<tr><td>1</td><td>PDPL official EN (Art.4–6, 10–15, 25–27)</td><td>SDAIA</td><td><a href="https://sdaia.gov.sa/en/SDAIA/about/Documents/Personal%20Data%20English%20V2-23April2023-%20Reviewed-.pdf">sdaia.gov.sa PDF</a></td></tr>
<tr><td>2</td><td>PDPL Art.26 direct marketing</td><td>SDAIA</td><td><a href="https://dgp.sdaia.gov.sa/wps/portal/pdp/knowledgecenter/details/PDPL2/">dgp.sdaia.gov.sa PDPL2</a></td></tr>
<tr><td>3</td><td>AI Ethics Principles v1.0</td><td>SDAIA</td><td><a href="https://dgp.sdaia.gov.sa/wps/wcm/connect/4c56ed1c-1b82-447d-ac29-638f5f99c12e/ai-principles-EN.pdf">ai-principles-EN.pdf</a></td></tr>
<tr><td>4</td><td>NSDAI (Vision 2030)</td><td>SDAIA</td><td><a href="https://sdaia.gov.sa/en/SDAIA/SdaiaStrategies/Documents/NSDAI.pdf">NSDAI.pdf</a></td></tr>
<tr><td>5</td><td>SERA residential consumption tariff</td><td>SERA</td><td><a href="https://www.sera.gov.sa/en/consumer/electric-tariff/electric-tariff-categories/consumption-tariff">sera.gov.sa tariff</a></td></tr>
<tr><td>6</td><td>SEC data sharing</td><td>SEC</td><td><a href="https://www.se.com.sa/en/Open-Data/Data-Sharing/">se.com.sa Open Data</a></td></tr>
<tr><td>7</td><td>KBEAT — Building Energy Assessment Tool</td><td>KAPSARC</td><td><a href="https://apps.kapsarc.org/appboard/kbeat/en">apps.kapsarc.org/kbeat</a></td></tr>
<tr><td>8</td><td>Financial Consumer Protection Principles (analog)</td><td>SAMA</td><td><a href="https://rulebook.sama.gov.sa/en/financial-consumer-protection-principles-and-rules">rulebook.sama.gov.sa</a></td></tr>
<tr><td>9</td><td>ITU-T Y.3172 ML pipeline architecture</td><td>ITU</td><td><a href="https://www.itu.int/rec/T-REC-Y.3172-201906-I">itu.int Y.3172</a></td></tr>
<tr><td>10</td><td>ISO/IEC 42001:2023 (benchmark only)</td><td>ISO</td><td><a href="https://www.iso.org/standard/81230.html">iso.org 81230</a></td></tr>
</tbody>
</table>
<table>
<thead><tr><th>Deliverable</th><th>Location</th></tr></thead>
<tbody>
<tr><td>GitHub</td><td><a href="https://github.com/aboalrejal-ai/wafier">https://github.com/aboalrejal-ai/wafier</a></td></tr>
<tr><td>Live demo</td><td><a href="https://wafier.aboalrejal.com/">https://wafier.aboalrejal.com/</a></td></tr>
<tr><td>Demo video</td><td><a href="https://www.youtube.com/watch?v=bRYwjbxs9t4">https://www.youtube.com/watch?v=bRYwjbxs9t4</a></td></tr>
<tr><td>Knowledge base</td><td><code>knowledge-base.json</code></td></tr>
<tr><td>Scenarios</td><td><code>scenarios/</code></td></tr>
<tr><td>Reproducible demo</td><td><code>pnpm demo</code></td></tr>
</tbody>
</table>
<p>Distributor: in-app + Web Notification API + Capacitor Local Notifications.</p>
</body>
</html>
"""
    HTML_PATH.write_text(html, encoding="utf-8")
    print("wrote", HTML_PATH)


if __name__ == "__main__":
    draw_pipeline(DIAGRAM)
    print("wrote", DIAGRAM)
    build_html()
    build_docx()
